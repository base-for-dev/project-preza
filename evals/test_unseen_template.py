"""The pipeline on templates nobody wrote code for.

The hackathon judges a template the team has never seen, so this builds two
templates from scratch — different slide sizes, fonts, palettes, layout names in
other languages, one with no placeholders at all — and runs everything that
needs no model (parse -> design system -> compose three variants -> audit ->
export) on each. If any step keyed off a sample template's name, size or
wording, one of these fails.

uv run pytest evals/test_unseen_template.py
"""

from __future__ import annotations

from pathlib import Path

import pytest
from audit import run_checks
from design_system import extract_design_system, mark_visual_titles
from export import export_pptx
from export.validate import validate
from generator.content import ChartSpec, DeckContent, DiagramSpec, SlideContent
from ir_schema import ChartSeries, DiagramItem
from layout import compose_deck
from parser import parse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

VARIANTS = ("compact", "standard", "detailed")


def _rename_layouts(presentation, names: dict[int, str]) -> None:
    for index, name in names.items():
        presentation.slide_layouts[index]._element.cSld.set("name", name)


def _style(paragraph_or_frame, font: str, size: float, color: str, bold=False) -> None:
    for run in [r for p in paragraph_or_frame.paragraphs for r in p.runs]:
        run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)


def _textbox(slide, left, top, width, height, text, font, size, color, bold=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    box.text_frame.text = text
    _style(box.text_frame, font, size, color, bold)
    return box


def _georgia_43(path: Path) -> Path:
    """4:3, serif type, warm palette, placeholders under Cyrillic layout names."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    _rename_layouts(prs, {0: "Обложка", 1: "Раздел с текстом", 5: "Только заголовок"})
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "Название доклада"
    cover.placeholders[1].text = "Подзаголовок"
    for _ in range(3):
        body = prs.slides.add_slide(prs.slide_layouts[1])
        body.shapes.title.text = "Заголовок раздела"
        body.placeholders[1].text = "Первый тезис\nВторой тезис\nТретий тезис"
        for shape in (body.shapes.title, body.placeholders[1]):
            _style(shape.text_frame, "Georgia", 28 if shape is body.shapes.title else 18, "5B2A0F")
    end = prs.slides.add_slide(prs.slide_layouts[5])
    end.shapes.title.text = "Спасибо за внимание"
    for slide in prs.slides:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FBF3E4")
    prs.save(str(path))
    return path


def _no_placeholders_169(path: Path) -> Path:
    """16:9, sans type, dark palette, every shape a plain text box named in another script."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12_192_000), Emu(6_858_000)
    blank = prs.slide_layouts[6]
    _rename_layouts(prs, {6: "空白"})
    for kind in ("cover", "cards", "cards", "text", "text", "closing"):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("101820")
        title = _textbox(
            slide, 0.8, 0.6, 11.5, 1.2, "Заголовок слайда", "Verdana", 40, "F2AA4C", True
        )
        title.name = "标题"
        if kind == "cover":
            _textbox(slide, 0.8, 2.4, 9, 1, "Подзаголовок доклада", "Verdana", 22, "FFFFFF")
        elif kind == "cards":
            for i in range(3):
                card = _textbox(
                    slide,
                    0.8 + i * 4,
                    2.6,
                    3.6,
                    3,
                    "Карточка текст пример",
                    "Verdana",
                    20,
                    "FFFFFF",
                )
                card.name = f"بطاقة {i}"
        elif kind == "text":
            _textbox(slide, 0.8, 2.2, 7.5, 4, "Основной текст слайда " * 4, "Verdana", 20, "FFFFFF")
    prs.save(str(path))
    return path


TEMPLATES = {"georgia-43": _georgia_43, "no-placeholders-169": _no_placeholders_169}


def _content(roles: list[str]) -> DeckContent:
    """Model-free content for `roles`: text everywhere, a diagram and a chart where they fit."""
    slides = []
    for i, role in enumerate(roles):
        slide = SlideContent(
            role=role,
            title=f"Вывод номер {i + 1} о состоянии проекта",
            bullets=["Первый пункт", "Второй пункт", "Третий пункт"],
            speaker_notes="Мы объясняем вывод слайда своими словами и переходим к следующему.",
        )
        if i == 2:
            slide.bullets = []
            slide.diagram = DiagramSpec(
                diagram_type="process",
                items=[
                    DiagramItem(label=f"Этап {n}", detail="Пояснение", icon="rocket")
                    for n in (1, 2, 3)
                ],
            )
        if i == 3:
            slide.chart = ChartSpec(
                title="Рост",
                unit="млн",
                category_label="Год",
                categories=["2024", "2025"],
                series=[ChartSeries(name="Выручка", values=[10, 20])],
            )
        slides.append(slide)
    return DeckContent(slides=slides)


@pytest.fixture(scope="module", params=sorted(TEMPLATES))
def unseen(request, tmp_path_factory):
    path = TEMPLATES[request.param](tmp_path_factory.mktemp(request.param) / "template.pptx")
    deck = mark_visual_titles(parse(path))
    design = extract_design_system(deck)
    return path, deck, design


def test_the_design_system_is_read_from_the_file(unseen):
    _, deck, design = unseen
    assert design.patterns, "no layouts found"
    assert design.typography.fonts, "no fonts found"
    assert {p.layout_name for p in design.patterns} == {s.layout_name for s in deck.slides}


@pytest.mark.parametrize("variant", VARIANTS)
def test_a_deck_is_composed_audited_and_exported_natively(unseen, variant, tmp_path):
    path, template_deck, design = unseen
    roles = [design.patterns[i % len(design.patterns)].layout_name for i in range(10)]
    composed = compose_deck(_content(roles), template_deck, variant)
    assert len(composed.slides) == 10

    findings = run_checks(composed, template_deck)
    broken = {"file_not_openable", "layout_not_from_template", "slide_is_picture"}
    assert not [f for f in findings if f.check in broken]

    out = tmp_path / "deck.pptx"
    export_pptx(composed, out, template_path=path)
    assert validate(out) == []
    reopened = Presentation(str(out))
    assert len(reopened.slides) == 10
    texts = [s.text_frame.text for sl in reopened.slides for s in sl.shapes if s.has_text_frame]
    assert any("Вывод номер" in t for t in texts), "titles are not native text"
    if variant == "standard":  # the other variants may put that slide on a design with no room
        assert any(getattr(sh, "has_chart", False) for sl in reopened.slides for sh in sl.shapes), (
            "the drawn chart is not a native chart in the file"
        )
    assert all(  # nothing is a picture of a slide
        not (len(sl.shapes) == 1 and sl.shapes[0].shape_type == 13) for sl in reopened.slides
    )


def test_speaker_notes_reach_the_exported_file(unseen, tmp_path):
    path, template_deck, design = unseen
    roles = [design.patterns[0].layout_name] * 3
    composed = compose_deck(_content(roles), template_deck, "standard")
    out = tmp_path / "deck.pptx"
    export_pptx(composed, out, template_path=path)
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(out)).slides]
    assert all("своими словами" in n for n in notes)
