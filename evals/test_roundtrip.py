"""Pytest coverage for the parse -> IR -> export round-trip.

    uv run pytest evals
"""

from __future__ import annotations

from pathlib import Path

import pytest
from export import export_pptx
from ir_schema import PassthroughShape, Table, TextBoxShape
from parser import parse
from pptx import Presentation

TEMPLATE_PATH = Path(__file__).parent / "templates" / "portrait-regiona.pptx"
OUTPUT_DIR = Path(__file__).parent / "output"


@pytest.fixture(scope="module")
def original_deck():
    return parse(TEMPLATE_PATH)


@pytest.fixture(scope="module")
def roundtrip_path(original_deck):
    out_path = OUTPUT_DIR / "portrait-regiona-test-roundtrip.pptx"
    export_pptx(original_deck, out_path)
    return out_path


@pytest.fixture(scope="module")
def roundtrip_deck(roundtrip_path):
    return parse(roundtrip_path)


def test_parse_does_not_crash():
    deck = parse(TEMPLATE_PATH)
    assert deck.slides


def test_parse_finds_all_slides(original_deck):
    assert len(original_deck.slides) == 15


def test_slide_count_roundtrips(original_deck, roundtrip_deck):
    assert len(roundtrip_deck.slides) == len(original_deck.slides)


def test_shape_count_roundtrips_excluding_passthrough(original_deck, roundtrip_deck):
    for original_slide, roundtrip_slide in zip(
        original_deck.slides, roundtrip_deck.slides, strict=True
    ):
        passthrough_count = sum(
            1 for shape in original_slide.shapes if isinstance(shape, PassthroughShape)
        )
        expected = len(original_slide.shapes) - passthrough_count
        assert len(roundtrip_slide.shapes) == expected, (
            f"slide {original_slide.index}: expected {expected}, "
            f"got {len(roundtrip_slide.shapes)}"
        )


def test_exported_pptx_is_valid(roundtrip_path):
    presentation = Presentation(str(roundtrip_path))
    assert len(presentation.slides) == 15


def test_text_box_content_survives_byte_for_byte(original_deck, roundtrip_deck):
    for original_slide, roundtrip_slide in zip(
        original_deck.slides, roundtrip_deck.slides, strict=True
    ):
        original_boxes = [s for s in original_slide.shapes if isinstance(s, TextBoxShape)]
        roundtrip_boxes = [s for s in roundtrip_slide.shapes if isinstance(s, TextBoxShape)]
        assert len(original_boxes) == len(roundtrip_boxes)
        for original_box, roundtrip_box in zip(original_boxes, roundtrip_boxes, strict=True):
            original_text = "".join(
                run.text for paragraph in original_box.paragraphs for run in paragraph.runs
            )
            roundtrip_text = "".join(
                run.text for paragraph in roundtrip_box.paragraphs for run in paragraph.runs
            )
            assert original_text == roundtrip_text


def test_table_content_survives_byte_for_byte(original_deck, roundtrip_deck):
    for original_slide, roundtrip_slide in zip(
        original_deck.slides, roundtrip_deck.slides, strict=True
    ):
        original_tables = [s for s in original_slide.shapes if isinstance(s, Table)]
        roundtrip_tables = [s for s in roundtrip_slide.shapes if isinstance(s, Table)]
        assert len(original_tables) == len(roundtrip_tables)
        for original_table, roundtrip_table in zip(original_tables, roundtrip_tables, strict=True):
            original_cells = [cell.text for row in original_table.rows for cell in row]
            roundtrip_cells = [cell.text for row in roundtrip_table.rows for cell in row]
            assert original_cells == roundtrip_cells
