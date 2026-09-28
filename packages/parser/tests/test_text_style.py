"""Inherited text size/font resolution (parser.text_style).

uv run pytest packages/parser
"""

from parser.parser import parse
from pptx import Presentation
from pptx.util import Pt


def test_title_placeholder_inherits_size_and_theme_font(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide
    slide.shapes.title.text = "Hello"
    path = tmp_path / "d.pptx"
    prs.save(str(path))

    run = next(
        r
        for s in parse(path).slides[0].shapes
        for p in getattr(s, "paragraphs", [])
        for r in p.runs
        if r.text == "Hello"
    )
    # python-pptx's default template: 44pt title in the theme's major font.
    assert run.font_size_pt == 44.0
    assert run.font_name == "Calibri"


def test_explicit_run_properties_win(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Hello"
    run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(20)
    run.font.name = "Oswald"
    path = tmp_path / "d.pptx"
    prs.save(str(path))

    parsed = next(
        r
        for s in parse(path).slides[0].shapes
        for p in getattr(s, "paragraphs", [])
        for r in p.runs
    )
    assert (parsed.font_size_pt, parsed.font_name) == (20.0, "Oswald")
