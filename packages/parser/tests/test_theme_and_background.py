"""Tests for slide background resolution and theme color resolution.

uv run pytest packages/parser -v
"""

from __future__ import annotations

from pathlib import Path

from ir_schema import Color
from parser import parse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Emu

TEMPLATE_PATH = Path(__file__).parents[2].parent / "evals" / "templates" / "portrait-regiona.pptx"


def test_slide_with_explicit_rgb_background_resolves(tmp_path):
    path = tmp_path / "rgb_bg.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    presentation.save(str(path))

    deck = parse(path)

    assert deck.slides[0].background == Color(kind="rgb", rgb="123456")


def test_slide_inherits_theme_color_background_from_master_through_clrmap(tmp_path):
    path = tmp_path / "theme_bg.pptx"
    presentation = Presentation()
    master = presentation.slide_masters[0]
    master.background.fill.solid()
    master.background.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
    # slide and layout both leave their own background unset -> must inherit
    # all the way down to the master.
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(str(path))

    deck = parse(path)

    background = deck.slides[0].background
    assert background is not None
    assert background.kind == "theme"
    assert background.theme_color == "bg1"

    # "bg1" is a semantic slot: it only resolves to a real color by way of
    # the slide master's <p:clrMap>, which points it at one of the 12 raw
    # <a:clrScheme> slots. Confirm the full chain landed on a real hex value
    # that agrees with whatever raw slot clrMap actually maps bg1 to.
    clr_map = master.element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap"
    )
    mapped_raw_slot = clr_map.get("bg1")
    assert deck.theme_colors["bg1"] == deck.theme_colors[mapped_raw_slot]
    assert len(deck.theme_colors["bg1"]) == 6
    int(deck.theme_colors["bg1"], 16)  # raises if not valid hex


def test_deck_theme_colors_contains_expected_slots_for_real_template():
    deck = parse(TEMPLATE_PATH)

    expected_slots = {
        "dk1",
        "lt1",
        "dk2",
        "lt2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
        "bg1",
        "tx1",
        "bg2",
        "tx2",
    }
    assert expected_slots <= deck.theme_colors.keys()
    for hex_value in deck.theme_colors.values():
        assert len(hex_value) == 6
        int(hex_value, 16)  # raises if not valid hex

    # The concrete case motivating this ticket: the template's slide
    # background is a theme-scheme color, and it should now resolve to real
    # branding instead of leaving the renderer to fall back to flat white.
    background = deck.slides[0].background
    assert background is not None
    assert background.kind == "theme"
    assert deck.theme_colors[background.theme_color] is not None


def test_theme_color_run_produces_clean_slot_name(tmp_path):
    path = tmp_path / "theme_run.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(914400), Emu(914400))
    run = textbox.text_frame.paragraphs[0].add_run()
    run.text = "hello"
    run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    presentation.save(str(path))

    deck = parse(path)

    runs = [
        run
        for shape in deck.slides[0].shapes
        for paragraph in shape.paragraphs
        for run in paragraph.runs
    ]
    assert len(runs) == 1
    color = runs[0].color
    assert color is not None
    assert color.kind == "theme"
    # Previously this was the raw `str(MSO_THEME_COLOR member)`, e.g.
    # "ACCENT_1 (5)" -- now it must be the clean OOXML slot name.
    assert color.theme_color == "accent1"
