"""Group members, photo-filled shapes and background photos in the IR.

uv run pytest packages/parser
"""

import io

from ir_schema import BACKGROUND_SHAPE_ID, Picture, TextBoxShape
from lxml import etree
from parser.parser import parse
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches


def _png() -> io.BytesIO:
    out = io.BytesIO()
    Image.new("RGB", (8, 8), (10, 120, 200)).save(out, format="PNG")
    out.seek(0)
    return out


def _deck_with_group(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text_frame.text = "Add a main point"
    group.shapes.add_picture(_png(), Inches(4), Inches(1), Inches(2), Inches(2))
    path = tmp_path / "g.pptx"
    prs.save(str(path))
    return path, box.shape_id


def test_text_and_pictures_inside_groups_are_listed_with_their_group(tmp_path):
    path, box_id = _deck_with_group(tmp_path)
    shapes = parse(path).slides[0].shapes
    member = next(s for s in shapes if s.shape_id == box_id)
    assert isinstance(member, TextBoxShape) and member.group_id is not None
    assert member.paragraphs[0].runs[0].text == "Add a main point"
    assert any(isinstance(s, Picture) and s.group_id for s in shapes)


def test_group_members_get_slide_coordinates(tmp_path):
    path, box_id = _deck_with_group(tmp_path)
    prs = Presentation(str(path))
    group = prs.slides[0].shapes[0]
    # Move the group: its child-space stays, its slide position shifts.
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    xfrm.find(qn("a:off")).set("x", str(int(xfrm.find(qn("a:off")).get("x")) + Inches(3)))
    prs.save(str(path))
    member = next(s for s in parse(path).slides[0].shapes if s.shape_id == box_id)
    assert member.left == Inches(4)


def test_shape_filled_with_a_photo_is_a_picture(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(_png(), 0, 0, Inches(2), Inches(2))
    rid = pic._element.blipFill.find(qn("a:blip")).get(qn("r:embed"))
    oval = slide.shapes.add_shape(9, Inches(3), 0, Inches(2), Inches(2))  # oval
    sp_pr = oval._element.spPr
    for fill in sp_pr.findall(qn("a:solidFill")):
        sp_pr.remove(fill)
    blip_fill = etree.SubElement(sp_pr, qn("a:blipFill"))
    etree.SubElement(blip_fill, qn("a:blip")).set(qn("r:embed"), rid)
    etree.SubElement(etree.SubElement(blip_fill, qn("a:stretch")), qn("a:fillRect"))
    sp_pr.remove(blip_fill)
    sp_pr.insert(2, blip_fill)  # after xfrm + prstGeom
    path = tmp_path / "f.pptx"
    prs.save(str(path))
    parsed = next(s for s in parse(path).slides[0].shapes if s.shape_id == oval.shape_id)
    assert isinstance(parsed, Picture) and parsed.image_bytes_b64


def test_background_photo_is_a_full_slide_picture(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(_png(), 0, 0, Emu(10), Emu(10))
    rid = pic._element.blipFill.find(qn("a:blip")).get(qn("r:embed"))
    c_sld = slide.element.find(qn("p:cSld"))
    bg = etree.Element(qn("p:bg"))
    bg_pr = etree.SubElement(bg, qn("p:bgPr"))
    fill = etree.SubElement(bg_pr, qn("a:blipFill"))
    etree.SubElement(fill, qn("a:blip")).set(qn("r:embed"), rid)
    etree.SubElement(bg_pr, qn("a:effectLst"))
    c_sld.insert(0, bg)
    path = tmp_path / "b.pptx"
    prs.save(str(path))
    background = parse(path).slides[0].shapes[0]
    assert background.shape_id == BACKGROUND_SHAPE_ID and isinstance(background, Picture)
    assert background.is_background  # design: never a photo slot
    assert (background.width, background.height) == (prs.slide_width, prs.slide_height)


def test_full_slide_picture_is_background_small_one_is_not(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(_png(), 0, 0, prs.slide_width, prs.slide_height)
    slide.shapes.add_picture(_png(), Inches(1), Inches(1), Inches(2), Inches(2))
    path = tmp_path / "p.pptx"
    prs.save(str(path))
    full, small = [s for s in parse(path).slides[0].shapes if isinstance(s, Picture)]
    assert full.is_background and not small.is_background
