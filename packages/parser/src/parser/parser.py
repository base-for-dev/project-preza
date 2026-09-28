"""PPTX -> IR parsing.

`parse()` walks every slide's shape tree with `python-pptx` and builds a
`Deck`. Nothing here assumes a particular template: shape kinds are dispatched
on `python-pptx`'s own `MSO_SHAPE_TYPE`, and anything not explicitly handled
falls through to `PassthroughShape` with its raw XML preserved.
"""

from __future__ import annotations

import base64
from pathlib import Path

from ir_schema import (
    BACKGROUND_SHAPE_ID,
    AutoShape,
    Color,
    Deck,
    Paragraph,
    PassthroughShape,
    Picture,
    Slide,
    Table,
    TableCell,
    TextBoxShape,
    TextRun,
)
from lxml import etree
from pptx import Presentation
from pptx.dml.color import ColorFormat
from pptx.dml.fill import FillFormat
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.base import BaseShape
from pptx.text.text import Font, TextFrame

from parser.text_style import Inherited, ShapeTextStyle, TextStyles

# MSO_THEME_COLOR member name -> the OOXML <a:schemeClr val="..."> slot it
# corresponds to. NOT_THEME_COLOR/MIXED are intentionally absent: callers
# treat an unmapped name as unresolvable, same as any other lookup miss.
_THEME_COLOR_SLOT = {
    "DARK_1": "dk1",
    "LIGHT_1": "lt1",
    "DARK_2": "dk2",
    "LIGHT_2": "lt2",
    "ACCENT_1": "accent1",
    "ACCENT_2": "accent2",
    "ACCENT_3": "accent3",
    "ACCENT_4": "accent4",
    "ACCENT_5": "accent5",
    "ACCENT_6": "accent6",
    "HYPERLINK": "hlink",
    "FOLLOWED_HYPERLINK": "folHlink",
    "BACKGROUND_1": "bg1",
    "BACKGROUND_2": "bg2",
    "TEXT_1": "tx1",
    "TEXT_2": "tx2",
}

# The 12 slots a <a:clrScheme> defines directly.
_RAW_THEME_SLOTS = [
    "dk1",
    "lt1",
    "dk2",
    "lt2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
]
# The 4 semantic slots a slide master's <p:clrMap> indirects to one of the
# raw slots above (accent1-6/hlink/folHlink always map to themselves).
_SEMANTIC_THEME_SLOTS = ["bg1", "tx1", "bg2", "tx2"]

_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def parse(path: Path) -> Deck:
    presentation = Presentation(str(path))
    theme_colors = _resolve_deck_theme_colors(presentation)
    styles = TextStyles(presentation)
    size = (presentation.slide_width, presentation.slide_height)
    slides = [
        _parse_slide(index, slide, styles, size) for index, slide in enumerate(presentation.slides)
    ]
    return Deck(
        slide_width=presentation.slide_width,
        slide_height=presentation.slide_height,
        source_path=str(path),
        slides=slides,
        theme_colors=theme_colors,
    )


def _parse_slide(index: int, slide, styles: TextStyles, size: tuple[int, int]) -> Slide:
    # Templates carry stray trailing spaces in layout names ("Мокап телефона ");
    # the model drops them, so exact-match role lookups would otherwise fail.
    layout_name = slide.slide_layout.name.strip()
    shapes = []
    background = _background_picture(slide, size)
    if background is not None:
        shapes.append(background)
    for z_order, shape in enumerate(slide.shapes):
        shapes.append(
            _parse_shape(
                shape, z_order, styles.for_shape(shape, slide) if shape.has_text_frame else None
            )
        )
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            shapes += _group_members(shape, shape.shape_id, z_order, slide, styles, _IDENTITY)
    _mark_backgrounds(shapes, size)
    return Slide(
        index=index,
        layout_name=layout_name,
        shapes=shapes,
        background=_resolve_slide_background(slide),
    )


# (x, y, w, h) in a group's child space -> slide space.
def _IDENTITY(x: int, y: int, w: int, h: int) -> tuple[int, int, int, int]:
    return x, y, w, h


def _group_transform(group, outer):
    """Child-space -> slide-space mapping for `group` nested in `outer`."""
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    parts = {} if xfrm is None else {c.tag.split("}")[1]: c for c in xfrm}
    off, ext, ch_off, ch_ext = (parts.get(k) for k in ("off", "ext", "chOff", "chExt"))
    if None in (off, ext, ch_off, ch_ext):
        return outer
    ox, oy = int(off.get("x")), int(off.get("y"))
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    chx, chy = int(ch_off.get("x")), int(ch_off.get("y"))
    chcx, chcy = int(ch_ext.get("cx")) or 1, int(ch_ext.get("cy")) or 1
    sx, sy = cx / chcx, cy / chcy

    def transform(x: int, y: int, w: int, h: int) -> tuple[int, int, int, int]:
        return outer(
            round(ox + (x - chx) * sx), round(oy + (y - chy) * sy), round(w * sx), round(h * sy)
        )

    return transform


def _group_members(group, group_id: int, z_order: int, slide, styles: TextStyles, outer) -> list:
    """Text- and picture-bearing shapes inside `group` (recursively), slide-positioned.

    Template designers group a card's heading with its text, or a portrait
    with its frame; without these the pipeline never saw that text or those
    photos, and generation left the template's "Add a main point" and stock
    photos on the finished slide. Pure decoration inside a group isn't
    listed — it stays in the group's passthrough XML.
    """
    transform = _group_transform(group, outer)
    members = []
    for child in group.shapes:
        if child.shape_type == MSO_SHAPE_TYPE.GROUP:
            members += _group_members(child, group_id, z_order, slide, styles, transform)
            continue
        has_text = child.has_text_frame and child.text_frame.text.strip()
        if not (has_text or _is_picture_like(child)):
            continue
        style = styles.for_shape(child, slide) if child.has_text_frame else None
        parsed = _parse_shape(child, z_order, style)
        if isinstance(parsed, PassthroughShape):
            continue
        left, top, width, height = transform(
            child.left or 0, child.top or 0, child.width or 0, child.height or 0
        )
        members.append(
            parsed.model_copy(
                update={
                    "left": left,
                    "top": top,
                    "width": width,
                    "height": height,
                    "group_id": group_id,
                }
            )
        )
    return members


def _background_picture(slide, size: tuple[int, int]) -> Picture | None:
    """The slide's own background image (`<p:bg>` image fill) as a full-slide
    `Picture` with `BACKGROUND_SHAPE_ID` and `is_background` — part of the
    slide's structure, never replaced; None for a colour or inherited one."""
    bg = slide.element.find(f"{qn('p:cSld')}/{qn('p:bg')}")
    blip = next(bg.iter(qn("a:blip")), None) if bg is not None else None
    rid = blip.get(qn("r:embed")) if blip is not None else None
    if not rid:
        return None
    try:
        part = slide.part.related_part(rid)
    except KeyError:
        return None
    return Picture(
        shape_id=BACKGROUND_SHAPE_ID,
        name="Slide background",
        z_order=-1,
        left=0,
        top=0,
        width=size[0],
        height=size[1],
        image_bytes_b64=_b64encode(part.blob),
        content_type=part.content_type,
        filename=str(part.partname).rsplit("/", 1)[-1],
        is_background=True,
    )


# A picture covering at least this share of the slide is its backdrop.
_BACKGROUND_SHARE = 0.75


def _mark_backgrounds(shapes: list, size: tuple[int, int]) -> None:
    """Flag full-slide pictures as the slide's background (never photo slots)."""
    slide_area = size[0] * size[1]
    for shape in shapes:
        covers = shape.width * shape.height >= _BACKGROUND_SHARE * slide_area
        if isinstance(shape, Picture) and covers:
            shape.is_background = True


def _blip_fill(shape: BaseShape):
    """The `<a:blip>` of a shape whose own fill is a picture, or None."""
    sp_pr = shape._element.find(qn("p:spPr"))
    fill = sp_pr.find(qn("a:blipFill")) if sp_pr is not None else None
    return fill.find(qn("a:blip")) if fill is not None else None


def _is_picture_like(shape: BaseShape) -> bool:
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE or _blip_fill(shape) is not None


def _parse_filled_shape(shape: BaseShape, z_order: int) -> Picture | None:
    """A shape filled with a photo (a portrait in a circle, a cover image in a
    blob) as a `Picture`, so generation can swap the photo; None if its image
    can't be read."""
    blip = _blip_fill(shape)
    rid = blip.get(qn("r:embed")) if blip is not None else None
    if not rid:
        return None
    try:
        part = shape.part.related_part(rid)
    except KeyError:
        return None
    return Picture(
        **_shape_base_fields(shape, z_order),
        image_bytes_b64=_b64encode(part.blob),
        content_type=part.content_type,
        filename=str(part.partname).rsplit("/", 1)[-1],
    )


def _resolve_deck_theme_colors(presentation) -> dict[str, str]:
    """Resolve every theme color slot to a hex string for this deck.

    A deck can have multiple slide masters with genuinely different themes
    (observed in real templates, e.g. vk-education.pptx and vktech.pptx both
    have 2 masters with distinct theme parts). `Deck.theme_colors` is a
    single flat dict shared by the whole deck, so full per-master accuracy
    is out of scope here (per ticket instructions) -- we resolve against the
    *first* slide master's theme. This is exactly correct for every
    single-master template (the common case, including
    evals/templates/portrait-regiona.pptx, the template motivating this
    ticket) and a documented approximation for multi-master decks: a slide
    belonging to a later master may have `Color.theme_color` slots that
    resolve to the wrong hex via this dict. Making that fully correct would
    need a per-slide or per-master color mapping, which is a bigger change
    than this ticket calls for.
    """
    if not presentation.slide_masters:
        return {}
    master = presentation.slide_masters[0]

    theme_part = master.part.part_related_by(RT.THEME)
    theme_xml = etree.fromstring(theme_part.blob)
    clr_scheme = theme_xml.find(f".//{_A_NS}clrScheme")
    if clr_scheme is None:
        return {}

    raw: dict[str, str] = {}
    for slot in _RAW_THEME_SLOTS:
        slot_el = clr_scheme.find(f"{_A_NS}{slot}")
        if slot_el is None:
            continue
        hex_value = _extract_theme_color_hex(slot_el)
        if hex_value is not None:
            raw[slot] = hex_value

    resolved = dict(raw)
    clr_map = master.element.find(f".//{_P_NS}clrMap")
    if clr_map is not None:
        for semantic_slot in _SEMANTIC_THEME_SLOTS:
            target_slot = clr_map.get(semantic_slot)
            if target_slot and target_slot in raw:
                resolved[semantic_slot] = raw[target_slot]

    return resolved


def _extract_theme_color_hex(slot_element) -> str | None:
    srgb = slot_element.find(f"{_A_NS}srgbClr")
    if srgb is not None:
        val = srgb.get("val")
        if val:
            return val.upper()
    sys_color = slot_element.find(f"{_A_NS}sysClr")
    if sys_color is not None:
        last_color = sys_color.get("lastClr")
        if last_color:
            return last_color.upper()
    return None


def _resolve_slide_background(slide) -> Color | None:
    for source in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        color = _resolve_fill_color(source.background.fill)
        if color is not None:
            return color
    return None


def _resolve_fill_color(fill: FillFormat) -> Color | None:
    try:
        fill_type = fill.type
    except (ValueError, AttributeError, TypeError):
        return None
    if fill_type != MSO_FILL_TYPE.SOLID:
        return None
    try:
        fore_color = fill.fore_color
    except (ValueError, AttributeError, TypeError):
        return None
    return _parse_color(fore_color)


def _placeholder_fields(shape: BaseShape) -> dict:
    if not shape.is_placeholder:
        return {"is_placeholder": False}
    ph_format = shape.placeholder_format
    return {
        "is_placeholder": True,
        "placeholder_type": str(ph_format.type) if ph_format.type is not None else None,
        "placeholder_idx": ph_format.idx,
    }


def _shape_base_fields(shape: BaseShape, z_order: int) -> dict:
    return {
        "shape_id": shape.shape_id,
        "name": shape.name,
        "z_order": z_order,
        "left": shape.left if shape.left is not None else 0,
        "top": shape.top if shape.top is not None else 0,
        "width": shape.width if shape.width is not None else 0,
        "height": shape.height if shape.height is not None else 0,
        "rotation": shape.rotation or 0.0,
        **_placeholder_fields(shape),
    }


def _parse_shape(shape: BaseShape, z_order: int, style: ShapeTextStyle | None = None) -> object:
    shape_type = shape.shape_type

    if shape.has_table:
        return _parse_table(shape, z_order)
    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        return _parse_picture(shape, z_order)
    has_text = shape.has_text_frame and shape.text_frame.text.strip()
    if _blip_fill(shape) is not None and not has_text:
        filled = _parse_filled_shape(shape, z_order)
        if filled is not None:
            return filled
    if shape_type in (MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER):
        return _parse_text_box(shape, z_order, style)
    if shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return _parse_autoshape(shape, z_order, style)
    if (
        shape_type == MSO_SHAPE_TYPE.FREEFORM
        and shape.has_text_frame
        and shape.text_frame.text.strip()
    ):
        # A custom-drawn shape holding text (a funnel stage, a label on a
        # blob) is a text slot like any autoshape; as passthrough its sample
        # text went through generation untouched.
        return _parse_autoshape(shape, z_order, style)

    return _parse_passthrough(shape, z_order)


def _parse_text_box(
    shape: BaseShape, z_order: int, style: ShapeTextStyle | None = None
) -> TextBoxShape:
    return TextBoxShape(
        **_shape_base_fields(shape, z_order),
        paragraphs=_parse_text_frame(shape.text_frame, style) if shape.has_text_frame else [],
    )


def _parse_autoshape(
    shape: BaseShape, z_order: int, style: ShapeTextStyle | None = None
) -> AutoShape:
    autoshape_type = None
    try:
        if shape.auto_shape_type is not None:
            autoshape_type = str(shape.auto_shape_type)
    except (ValueError, AttributeError):
        pass

    fill_color = None
    try:
        if shape.fill.type is not None:
            fill_color = _parse_color(shape.fill.fore_color)
    except (ValueError, AttributeError, TypeError):
        pass

    return AutoShape(
        **_shape_base_fields(shape, z_order),
        autoshape_type=autoshape_type,
        fill_color=fill_color,
        paragraphs=_parse_text_frame(shape.text_frame, style) if shape.has_text_frame else [],
    )


def _parse_picture(shape: BaseShape, z_order: int) -> Picture:
    image = shape.image
    crop_left = shape.crop_left or 0.0
    crop_top = shape.crop_top or 0.0
    crop_right = shape.crop_right or 0.0
    crop_bottom = shape.crop_bottom or 0.0
    return Picture(
        **_shape_base_fields(shape, z_order),
        image_bytes_b64=_b64encode(image.blob),
        content_type=image.content_type,
        filename=image.filename,
        crop_left=crop_left,
        crop_top=crop_top,
        crop_right=crop_right,
        crop_bottom=crop_bottom,
    )


def _parse_table(shape: BaseShape, z_order: int) -> Table:
    table = shape.table
    rows = [
        [TableCell(paragraphs=_parse_text_frame(cell.text_frame)) for cell in row.cells]
        for row in table.rows
    ]
    return Table(
        **_shape_base_fields(shape, z_order),
        rows=rows,
        column_widths=[col.width for col in table.columns],
        row_heights=[row.height for row in table.rows],
    )


def _parse_passthrough(shape: BaseShape, z_order: int) -> PassthroughShape:
    return PassthroughShape(
        **_shape_base_fields(shape, z_order),
        original_shape_type=str(shape.shape_type) if shape.shape_type is not None else None,
        raw_xml=etree.tostring(shape._element, encoding="unicode"),
    )


def _parse_text_frame(
    text_frame: TextFrame, style: ShapeTextStyle | None = None
) -> list[Paragraph]:
    return [_parse_paragraph(paragraph, style) for paragraph in text_frame.paragraphs]


def _parse_paragraph(paragraph, style: ShapeTextStyle | None = None) -> Paragraph:
    # What the paragraph's runs inherit where they don't say (see text_style).
    inherited = style.level(paragraph.level or 0) if style is not None else Inherited()
    runs = [_style_run(run.text, run.font, inherited, style) for run in paragraph.runs]
    if not runs:
        # An empty slot (typically a placeholder awaiting text) has no runs,
        # but its end-of-paragraph properties say how typed text will look —
        # often 40pt+ on a title slide. Without them every capacity and
        # shrink-to-fit estimate falls back to a generic 18pt and badly
        # overestimates how much text fits. Kept as one empty-text run so
        # "has text" checks are unaffected.
        end = paragraph._p.find(qn("a:endParaRPr"))
        if end is not None and end.get("sz"):
            runs = [_style_run("", Font(end), inherited, style)]
    return Paragraph(
        runs=runs,
        alignment=str(paragraph.alignment) if paragraph.alignment is not None else None,
        level=paragraph.level or 0,
    )


def _style_run(
    text: str,
    font: Font,
    inherited: Inherited = Inherited(),  # noqa: B008 (frozen, never mutated)
    style: ShapeTextStyle | None = None,
) -> TextRun:
    # A run's own properties win; anything it leaves unsaid is what it
    # inherits from its placeholder / master / default text styles.
    name = style.font(font.name) if style is not None else font.name
    return TextRun(
        text=text,
        font_name=name or inherited.font_name,
        font_size_pt=font.size.pt if font.size is not None else inherited.size_pt,
        bold=font.bold if font.bold is not None else inherited.bold,
        italic=font.italic,
        underline=font.underline if isinstance(font.underline, bool) else None,
        color=_parse_color(font.color),
    )


def _parse_color(color_format: ColorFormat) -> Color | None:
    try:
        color_type = color_format.type
    except AttributeError:
        return None
    if color_type is None:
        return None

    if color_type.name == "RGB":
        return Color(kind="rgb", rgb=str(color_format.rgb))
    if color_type.name == "SCHEME":
        slot = _THEME_COLOR_SLOT.get(color_format.theme_color.name)
        if slot is None:
            return None
        return Color(kind="theme", theme_color=slot)
    return None


def _b64encode(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")
