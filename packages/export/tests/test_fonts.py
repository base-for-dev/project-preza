"""uv run pytest packages/export"""

import io
import re
import zipfile
from pathlib import Path

import pytest
from export import fonts
from export.mtx import _undo_run_length, eot_to_ttf
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Inches

TEMPLATES = Path(__file__).resolve().parents[3] / "evals" / "templates"


def _tiny_font(family: str = "Tiny", weight: int = 400) -> bytes:
    fb = FontBuilder(1000, isTTF=True)
    fb.setupGlyphOrder([".notdef", "A"])
    fb.setupCharacterMap({ord("A"): "A"})
    pen = TTGlyphPen(None)
    pen.moveTo((0, 0))
    pen.lineTo((500, 700))
    pen.lineTo((600, 0))
    pen.closePath()
    fb.setupGlyf({".notdef": TTGlyphPen(None).glyph(), "A": pen.glyph()})
    fb.setupHorizontalMetrics({".notdef": (500, 0), "A": (600, 0)})
    fb.setupHorizontalHeader(ascent=800, descent=-200)
    fb.setupNameTable({"familyName": family, "styleName": "Regular"})
    fb.setupOS2(usWeightClass=weight)
    fb.setupPost()
    out = io.BytesIO()
    fb.save(out)
    return out.getvalue()


def _deck(path: Path, font_name: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Hello"
    box.text_frame.paragraphs[0].runs[0].font.name = font_name
    prs.save(str(path))


@pytest.fixture
def cache(tmp_path, monkeypatch):
    monkeypatch.setattr(fonts, "CACHE_DIR", tmp_path / "cache")
    return tmp_path / "cache"


def test_split_weight_reads_google_slides_weight_families():
    assert fonts._split_weight("Montserrat SemiBold") == ("Montserrat", 600)
    assert fonts._split_weight("Fira Sans Extra Bold") == ("Fira Sans", 800)
    assert fonts._split_weight("Inter-Regular") == ("Inter", 400)
    assert fonts._split_weight("Lato") == ("Lato", 400)


def test_used_families_skips_system_fonts_and_theme_refs(tmp_path):
    path = tmp_path / "d.pptx"
    _deck(path, "Chakra Petch")
    families = fonts.used_families(path)
    assert "Chakra Petch" in families
    assert "Calibri" not in families  # python-pptx's default theme font
    assert not any(f.startswith("+") for f in families)


def test_uncompressed_eot_round_trips():
    ttf = _tiny_font()
    assert eot_to_ttf(fonts.to_eot(ttf)) == ttf


def test_run_length_escape_expands_runs_and_literal_escapes():
    # escape byte 0xAA; "AA 03 07" = three 0x07; "AA 00" = a literal 0xAA
    assert _undo_run_length(bytes([0xAA, 1, 0xAA, 3, 7, 2, 0xAA, 0])) == bytes(
        [1, 7, 7, 7, 2, 0xAA]
    )


def test_with_fonts_embeds_fetched_family_under_the_decks_name(tmp_path, cache, monkeypatch):
    src, dst = tmp_path / "in.pptx", tmp_path / "out.pptx"
    _deck(src, "Montserrat SemiBold")
    monkeypatch.setattr(
        fonts, "_google_fonts", lambda family: {"regular": _tiny_font("Montserrat", 600)}
    )

    assert fonts.with_fonts(src, dst)

    with zipfile.ZipFile(dst) as z:
        presentation = z.read("ppt/presentation.xml").decode()
        rels = z.read("ppt/_rels/presentation.xml.rels").decode()
        assert 'Extension="fntdata"' in z.read("[Content_Types].xml").decode()
        rid = re.search(
            r'typeface="Montserrat SemiBold"/><p:regular r:id="(rId\d+)"', presentation
        ).group(1)
        target = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"', rels).group(1)
        font = TTFont(io.BytesIO(eot_to_ttf(z.read(f"ppt/{target}"))))
    # Renamed to exactly what the deck asks for, so renderers match it.
    assert font["name"].getDebugName(1) == "Montserrat SemiBold"
    assert font["name"].getDebugName(2) == "Regular"
    assert fonts.font_file("Montserrat SemiBold", "regular") is not None
    # The source deck itself is untouched.
    with zipfile.ZipFile(src) as z:
        assert "embeddedFontLst" not in z.read("ppt/presentation.xml").decode()


def test_with_fonts_plain_copy_when_nothing_to_add(tmp_path, cache, monkeypatch):
    src, dst = tmp_path / "in.pptx", tmp_path / "out.pptx"
    _deck(src, "Arial")
    monkeypatch.setattr(fonts, "_google_fonts", lambda family: {})
    assert not fonts.with_fonts(src, dst)
    assert dst.read_bytes() == src.read_bytes()


def test_font_css_lists_only_cached_styles(cache, monkeypatch, tmp_path):
    src = tmp_path / "in.pptx"
    _deck(src, "Tiny")
    monkeypatch.setattr(fonts, "_google_fonts", lambda family: {"bold": _tiny_font("Tiny", 700)})
    fonts.template_fonts(src)
    css = fonts.font_css(["Tiny", "Nope"], lambda f, s: f"/f/{f}/{s}")
    assert css.count("@font-face") == 1
    assert "font-weight:700" in css and 'url("/f/Tiny/bold")' in css


@pytest.mark.skipif(
    not (TEMPLATES / "3d-modern-background-pitch-deck.pptx").is_file(),
    reason="sample template not present (evals/templates is gitignored)",
)
def test_decodes_mtx_compressed_embedded_font():
    with zipfile.ZipFile(TEMPLATES / "3d-modern-background-pitch-deck.pptx") as z:
        font = TTFont(io.BytesIO(eot_to_ttf(z.read("ppt/fonts/ChakraPetch-bold.fntdata"))))
    assert font["name"].getDebugName(4) == "Chakra Petch Bold"
    glyf = font["glyf"]
    a = glyf[font.getBestCmap()[ord("A")]]
    a.expand(glyf)
    assert a.numberOfContours > 0 and a.xMax > a.xMin


def test_char_width_of_a_font_is_its_average_advance():
    # _tiny_font maps only "A" (600/1000 em); everything else falls back.
    width = fonts.char_width_em(_tiny_font())
    assert 0.5 < width < 0.6


def test_annotate_char_widths_sets_runs_in_known_fonts(tmp_path, cache, monkeypatch):
    from parser.parser import parse

    src = tmp_path / "in.pptx"
    _deck(src, "Tiny")
    monkeypatch.setattr(fonts, "_google_fonts", lambda family: {"regular": _tiny_font("Tiny")})
    deck = parse(src)
    fonts.annotate_char_widths(deck, src)
    runs = [
        r
        for s in deck.slides
        for sh in s.shapes
        for p in getattr(sh, "paragraphs", [])
        for r in p.runs
    ]
    assert runs and all(r.char_width_em for r in runs if r.font_name == "Tiny")
