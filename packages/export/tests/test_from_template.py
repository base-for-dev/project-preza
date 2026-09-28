"""uv run pytest packages/export"""

from export.export import export_pptx
from lxml import etree
from parser.parser import parse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches


def _template(path):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for label in ("Cover", "Body"):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0x44, 0x11, 0x88)
        keep = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        keep.text_frame.text = f"{label} heading"
        keep.text_frame.paragraphs[0].runs[0].font.size = Inches(0.4)
        drop = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
        drop.text_frame.text = "Имя Фамилия"
    prs.save(str(path))


def test_export_clones_source_slide_keeps_design_and_swaps_content(tmp_path):
    template = tmp_path / "t.pptx"
    _template(template)
    deck = parse(template)

    composed = deck.model_copy(deep=True)
    slide = composed.slides[1].model_copy(deep=True)  # built from template slide #2
    slide.index, slide.source_index, slide.notes = 0, 1, "Говорим"
    heading, sample = slide.shapes
    heading.paragraphs[0].runs[0].text = "Новый заголовок"
    slide.shapes = [heading]  # composition removed the sample text box
    composed.slides = [slide]

    out = tmp_path / "out.pptx"
    export_pptx(composed, out, template_path=template)

    result = Presentation(str(out))
    assert len(result.slides) == 1  # template's own slides dropped
    new = result.slides[0]
    texts = [s.text_frame.text for s in new.shapes if s.has_text_frame]
    assert texts == ["Новый заголовок"]
    # Template formatting (run size) and the slide background survive.
    run = new.shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size == Inches(0.4)
    assert new.element.find(qn("p:cSld")).find(qn("p:bg")) is not None
    assert new.notes_slide.notes_text_frame.text == "Говорим"


def test_empty_paragraph_keeps_its_end_font_size(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    p = box.text_frame.paragraphs[0]._p
    end = etree.SubElement(p, qn("a:endParaRPr"))
    end.set("sz", "4000")
    path = tmp_path / "e.pptx"
    prs.save(str(path))

    shape = parse(path).slides[0].shapes[0]

    run = shape.paragraphs[0].runs[0]
    assert (run.text, run.font_size_pt) == ("", 40.0)


def test_chart_gets_composed_data_and_clones_do_not_share_it(tmp_path):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sample = CategoryChartData()
    sample.categories = ["2021", "2022"]
    sample.add_series("Ряд 1", [4.3, 2.5])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), sample)
    template = tmp_path / "c.pptx"
    prs.save(str(template))

    deck = parse(template)
    first, second = (deck.slides[0].model_copy(deep=True) for _ in range(2))
    for i, (s, value) in enumerate(((first, "10"), (second, "20"))):
        s.index, s.source_index = i, 0
        chart = s.shapes[0]
        assert chart.is_chart
        chart.chart_data = [["Метрика", "Секунды"], ["Генерация", value]]
    composed = deck.model_copy(update={"slides": [first, second]})

    out = tmp_path / "out.pptx"
    export_pptx(composed, out, template_path=template)

    charts = [s.shapes[0].chart for s in Presentation(str(out)).slides]
    assert [list(c.plots[0].categories) for c in charts] == [["Генерация"], ["Генерация"]]
    assert [c.series[0].values for c in charts] == [(10.0,), (20.0,)]
    assert charts[0].series[0].name == "Секунды"


def test_image_goes_into_empty_picture_placeholder(tmp_path):
    import base64
    import io as _io

    from ir_schema import Picture
    from PIL import Image

    prs = Presentation()
    layout = next(
        lo
        for lo in prs.slide_layouts
        if any(p.placeholder_format.type == 18 for p in lo.placeholders)
    )
    prs.slides.add_slide(layout)
    template = tmp_path / "p.pptx"
    prs.save(str(template))

    deck = parse(template)
    slide = deck.slides[0].model_copy(deep=True)
    slide.source_index = 0
    frame = next(s for s in slide.shapes if "PICTURE" in (s.placeholder_type or ""))
    png = _io.BytesIO()
    Image.new("RGB", (40, 30), "red").save(png, format="PNG")
    picture = Picture(
        **frame.model_dump(exclude={"kind", "paragraphs"}),
        image_bytes_b64=base64.b64encode(png.getvalue()).decode(),
        content_type="image/png",
    )
    slide.shapes = [picture if s is frame else s for s in slide.shapes]

    out = tmp_path / "out.pptx"
    export_pptx(deck.model_copy(update={"slides": [slide]}), out, template_path=template)

    shapes = Presentation(str(out)).slides[0].shapes
    assert any(
        s.shape_type == 13 or getattr(s, "image", None) is not None
        for s in shapes
        if s.shape_id == frame.shape_id
    )


def test_group_members_are_written_and_dropped_ones_blanked(tmp_path):
    from export.from_template import export_from_template
    from parser.parser import parse as _parse
    from pptx import Presentation as _P
    from pptx.util import Inches as _In

    prs = _P()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    keep = group.shapes.add_textbox(_In(1), _In(1), _In(2), _In(1))
    keep.text_frame.text = "Add a main point"
    drop = group.shapes.add_textbox(_In(1), _In(2), _In(2), _In(1))
    drop.text_frame.text = "Elaborate here"
    template = tmp_path / "t.pptx"
    prs.save(str(template))

    deck = _parse(template)
    deck.slides[0].source_index = 0
    member = next(s for s in deck.slides[0].shapes if s.shape_id == keep.shape_id)
    member.paragraphs[0].runs[0].text = "Новый пункт"
    deck.slides[0].shapes = [s for s in deck.slides[0].shapes if s.shape_id != drop.shape_id]
    out = tmp_path / "o.pptx"
    export_from_template(deck, template, out)

    texts = [sh.text_frame.text for sh in _P(str(out)).slides[0].shapes[0].shapes]
    assert texts == ["Новый пункт", ""]
