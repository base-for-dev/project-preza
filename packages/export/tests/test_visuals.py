"""Drawn charts, diagrams and pictograms come out as native, editable objects."""

from __future__ import annotations

import pytest
from export import export_pptx
from export.icons import ICONS
from export.validate import validate
from ir_schema import (
    ICON_NAMES,
    ChartSeries,
    ChartShape,
    Color,
    Deck,
    DiagramItem,
    DiagramShape,
    Slide,
)
from pptx import Presentation
from pptx.util import Emu

W, H = 12_192_000, 6_858_000
THEME = {"accent1": "0077FF", "accent2": "FF6B00", "accent3": "1DB954", "accent4": "8E44AD"}
BOX = {"left": 800_000, "top": 1_600_000, "width": 10_500_000, "height": 4_400_000}


def _deck(*shapes) -> Deck:
    slide = Slide(
        index=0,
        layout_name="L",
        shapes=list(shapes),
        background=Color(kind="rgb", rgb="FFFFFF"),
    )
    return Deck(slide_width=W, slide_height=H, slides=[slide], theme_colors=THEME)


def _export(deck: Deck, tmp_path):
    out = tmp_path / "deck.pptx"
    export_pptx(deck, out)
    assert validate(out) == []
    return Presentation(str(out))


def _items(n: int, icons: bool = False) -> list[DiagramItem]:
    names = list(ICON_NAMES)
    return [
        DiagramItem(label=f"Шаг {i + 1}", detail="Пояснение к шагу", icon=names[i] if icons else "")
        for i in range(n)
    ]


def _diagram(kind: str, n: int, icons: bool = False) -> DiagramShape:
    return DiagramShape(
        shape_id=10, name="d", z_order=5, diagram_type=kind, items=_items(n, icons), **BOX
    )


@pytest.mark.parametrize("chart_type", ["column", "bar", "line", "pie", "doughnut"])
def test_a_chart_is_a_real_chart_with_its_data(chart_type, tmp_path):
    chart = ChartShape(
        shape_id=9,
        name="c",
        z_order=5,
        chart_type=chart_type,
        categories=["2023", "2024", "2025"],
        series=[ChartSeries(name="Выручка", values=[10, 24, 51])],
        unit="млн ₽",
        category_label="Год",
        **BOX,
    )
    slide = _export(_deck(chart), tmp_path).slides[0]
    frames = [s for s in slide.shapes if getattr(s, "has_chart", False)]
    assert len(frames) == 1
    plot = frames[0].chart.plots[0]
    assert list(plot.categories) == ["2023", "2024", "2025"]
    assert list(plot.series[0].values) == [10, 24, 51]


def test_a_bar_chart_gets_axis_titles_and_a_legend_for_several_series(tmp_path):
    chart = ChartShape(
        shape_id=9,
        name="c",
        z_order=5,
        categories=["Q1", "Q2"],
        series=[ChartSeries(name="Мы", values=[1, 2]), ChartSeries(name="Рынок", values=[2, 3])],
        unit="%",
        category_label="Квартал",
        **BOX,
    )
    chart_part = next(
        s for s in _export(_deck(chart), tmp_path).slides[0].shapes if s.has_chart
    ).chart
    assert chart_part.has_legend
    assert chart_part.value_axis.axis_title.text_frame.text == "%"
    assert chart_part.category_axis.axis_title.text_frame.text == "Квартал"


@pytest.mark.parametrize(
    ("kind", "count"),
    [("process", 4), ("timeline", 5), ("cycle", 4), ("hierarchy", 5), ("icons", 6)],
)
def test_a_diagram_is_one_group_of_native_shapes_with_the_labels_as_text(kind, count, tmp_path):
    slide = _export(_deck(_diagram(kind, count)), tmp_path).slides[0]
    groups = [s for s in slide.shapes if s.shape_type == 6]  # MSO_SHAPE_TYPE.GROUP
    assert len(groups) == 1
    texts = " ".join(s.text_frame.text for s in groups[0].shapes if s.has_text_frame)
    for i in range(count):
        assert f"Шаг {i + 1}" in texts


def test_every_pictogram_name_has_a_drawing_and_draws(tmp_path):
    assert set(ICONS) == set(ICON_NAMES)
    shape = _diagram("icons", len(ICON_NAMES), icons=True)
    shape.items = [DiagramItem(label=n, icon=n) for n in ICON_NAMES]
    # More than the diagram keeps: draw them in chunks so each icon is exercised.
    for start in range(0, len(ICON_NAMES), 8):
        chunk = shape.model_copy(update={"items": shape.items[start : start + 8]})
        _export(_deck(chunk), tmp_path)


def test_a_diagram_with_no_items_or_a_chart_with_no_data_draws_nothing(tmp_path):
    empty_chart = ChartShape(shape_id=9, name="c", z_order=5, **BOX)
    empty_diagram = _diagram("process", 0)
    slide = _export(_deck(empty_chart, empty_diagram), tmp_path).slides[0]
    assert len(slide.shapes) == 0


def test_drawn_shapes_stay_inside_their_box(tmp_path):
    slide = _export(_deck(_diagram("process", 6, icons=True)), tmp_path).slides[0]
    group = next(s for s in slide.shapes if s.shape_type == 6)
    assert group.left >= Emu(BOX["left"]) - 1
    assert group.top >= Emu(BOX["top"]) - 1
    assert group.left + group.width <= Emu(BOX["left"] + BOX["width"]) + 1
