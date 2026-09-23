"""uv run pytest packages/export"""

from export.export import export_pptx
from ir_schema import Deck, Slide
from pptx import Presentation


def test_notes_written_to_notes_page(tmp_path):
    deck = Deck(
        slide_width=9_144_000,
        slide_height=6_858_000,
        slides=[
            Slide(index=0, layout_name="L", notes="Текст выступления"),
            Slide(index=1, layout_name="L"),
        ],
    )
    out = tmp_path / "deck.pptx"

    export_pptx(deck, out)

    slides = Presentation(str(out)).slides
    assert slides[0].notes_slide.notes_text_frame.text == "Текст выступления"
    assert not slides[1].has_notes_slide
