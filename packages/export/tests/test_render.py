"""uv run pytest packages/export"""

import pytest
from export import render
from export.render import RenderUnavailable, render_pptx, soffice_path
from pptx import Presentation


def test_soffice_path_prefers_configured(monkeypatch, tmp_path):
    fake = tmp_path / "soffice"
    fake.write_text("")
    monkeypatch.setenv("SOFFICE_PATH", str(fake))
    assert soffice_path() == str(fake)


def test_missing_libreoffice_raises_render_unavailable(monkeypatch, tmp_path):
    monkeypatch.setattr(render, "soffice_path", lambda: None)
    with pytest.raises(RenderUnavailable):
        render_pptx(tmp_path / "x.pptx")


@pytest.mark.skipif(soffice_path() is None, reason="LibreOffice not installed")
def test_renders_one_png_per_slide(tmp_path):
    prs = Presentation()
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "d.pptx"
    prs.save(str(path))

    pages = render_pptx(path, width_px=320)

    assert len(pages) == 2
    assert all(p.startswith(b"\x89PNG") for p in pages)
