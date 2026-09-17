"""IR -> PPTX export.

`export_pptx()` reconstructs a `.pptx` from a `Deck` using native, editable
`python-pptx` objects — text boxes with runs, autoshapes, tables. Shapes
don't get remapped onto template layouts (that's `packages/layout`'s job);
each slide is built on a blank layout with shapes placed at their IR
coordinates, which is enough to prove text/shape/table fidelity round-trips.

`PassthroughShape` (charts, SmartArt, groups, ...) is intentionally skipped
for this PoC — logged, not silently dropped, and not rasterized.
"""

from __future__ import annotations

import base64
import io
import logging
from pathlib import Path

from ir_schema import (
    AutoShape,
    Color,
    Deck,
    Paragraph,
    PassthroughShape,
    Picture,
    Shape,
    Slide,
    Table,
    TextBoxShape,
    TextRun,
)
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

logger = logging.getLogger(__name__)

_ALIGNMENT_BY_NAME = {member.name: member for member in PP_ALIGN}


def export_pptx(deck: Deck, out_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Emu(deck.slide_width)
    presentation.slide_height = Emu(deck.slide_height)
    blank_layout = _find_blank_layout(presentation)

    for slide_ir in deck.slides:
        slide = presentation.slides.add_slide(blank_layout)
        _export_slide(slide_ir, slide)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(out_path))


def _find_blank_layout(presentation: Presentation):
    for layout in presentation.slide_layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return presentation.slide_layouts[-1]


def _export_slide(slide_ir: Slide, slide) -> None:
    for shape_ir in sorted(slide_ir.shapes, key=lambda shape: shape.z_order):
        _export_shape(shape_ir, slide)


def _export_shape(shape_ir: Shape, slide) -> None:
    if isinstance(shape_ir, TextBoxShape):
        _export_text_box(shape_ir, slide)
    elif isinstance(shape_ir, AutoShape):
        _export_autoshape(shape_ir, slide)
    elif isinstance(shape_ir, Table):
        _export_table(shape_ir, slide)
    elif isinstance(shape_ir, Picture):
        _export_picture(shape_ir, slide)
    elif isinstance(shape_ir, PassthroughShape):
        logger.info(
            "skipping passthrough shape %r (%s) on export",
            shape_ir.name,
            shape_ir.original_shape_type,
        )
    else:  # pragma: no cover - exhaustive over the Shape union
        raise TypeError(f"unhandled shape kind: {type(shape_ir)!r}")


def _export_text_box(shape_ir: TextBoxShape, slide) -> None:
    box = slide.shapes.add_textbox(
        Emu(shape_ir.left), Emu(shape_ir.top), Emu(shape_ir.width), Emu(shape_ir.height)
    )
    box.rotation = shape_ir.rotation
    _fill_text_frame(box.text_frame, shape_ir.paragraphs)


def _export_autoshape(shape_ir: AutoShape, slide) -> None:
    shape_type = _resolve_autoshape_type(shape_ir.autoshape_type) or MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Emu(shape_ir.left), Emu(shape_ir.top), Emu(shape_ir.width), Emu(shape_ir.height)
    )
    shape.rotation = shape_ir.rotation
    if shape_ir.fill_color is not None:
        _apply_fill_color(shape, shape_ir.fill_color)
    _fill_text_frame(shape.text_frame, shape_ir.paragraphs)


def _resolve_autoshape_type(name: str | None):
    if not name:
        return None
    bare_name = name.split(" (")[0]
    try:
        return MSO_SHAPE[bare_name]
    except KeyError:
        return None


def _apply_fill_color(shape, color: Color) -> None:
    if color.kind != "rgb" or not color.rgb:
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color.rgb)


def _export_table(shape_ir: Table, slide) -> None:
    n_rows = len(shape_ir.rows)
    n_cols = len(shape_ir.rows[0]) if n_rows else 0
    if n_rows == 0 or n_cols == 0:
        logger.info("skipping empty table shape %r on export", shape_ir.name)
        return

    graphic_frame = slide.shapes.add_table(
        n_rows,
        n_cols,
        Emu(shape_ir.left),
        Emu(shape_ir.top),
        Emu(shape_ir.width),
        Emu(shape_ir.height),
    )
    table = graphic_frame.table

    for col_idx, width in enumerate(shape_ir.column_widths):
        if col_idx < len(table.columns):
            table.columns[col_idx].width = Emu(width)
    for row_idx, height in enumerate(shape_ir.row_heights):
        if row_idx < len(table.rows):
            table.rows[row_idx].height = Emu(height)

    for row_idx, row in enumerate(shape_ir.rows):
        for col_idx, cell_ir in enumerate(row):
            _fill_text_frame(table.cell(row_idx, col_idx).text_frame, cell_ir.paragraphs)


def _export_picture(shape_ir: Picture, slide) -> None:
    if not shape_ir.image_bytes_b64:
        logger.info("skipping picture shape %r with no captured bytes", shape_ir.name)
        return
    image_stream = io.BytesIO(base64.b64decode(shape_ir.image_bytes_b64))
    slide.shapes.add_picture(
        image_stream,
        Emu(shape_ir.left),
        Emu(shape_ir.top),
        Emu(shape_ir.width),
        Emu(shape_ir.height),
    )


def _fill_text_frame(text_frame, paragraphs: list[Paragraph]) -> None:
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.TOP

    for index, paragraph_ir in enumerate(paragraphs):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        _fill_paragraph(paragraph, paragraph_ir)


def _fill_paragraph(paragraph, paragraph_ir: Paragraph) -> None:
    paragraph.level = paragraph_ir.level
    if paragraph_ir.alignment:
        bare_name = paragraph_ir.alignment.split(" (")[0]
        paragraph.alignment = _ALIGNMENT_BY_NAME.get(bare_name)

    for run_ir in paragraph_ir.runs:
        run = paragraph.add_run()
        _fill_run(run, run_ir)


def _fill_run(run, run_ir: TextRun) -> None:
    run.text = run_ir.text
    font = run.font
    if run_ir.font_name is not None:
        font.name = run_ir.font_name
    if run_ir.font_size_pt is not None:
        font.size = Pt(run_ir.font_size_pt)
    if run_ir.bold is not None:
        font.bold = run_ir.bold
    if run_ir.italic is not None:
        font.italic = run_ir.italic
    if run_ir.underline is not None:
        font.underline = run_ir.underline
    if run_ir.color is not None and run_ir.color.kind == "rgb" and run_ir.color.rgb:
        font.color.rgb = RGBColor.from_string(run_ir.color.rgb)
