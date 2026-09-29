"""Charts and diagrams the pipeline draws itself, as native PowerPoint objects.

`ChartShape` becomes a real embedded chart (data editable in PowerPoint);
`DiagramShape` becomes a group of ordinary shapes, text boxes and arrows — the
editable stand-in for SmartArt. Colours come from the deck's theme accents,
the font from the template's main font, so a drawn visual looks like the
template's own.
"""

from __future__ import annotations

import math

from ir_schema import ChartShape, Deck, DiagramItem, DiagramShape, Slide
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from export.icons import draw_icon

_EMU_PT = 12700
_CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}
DEFAULT_TEXT_PT = 14.0
MIN_TEXT_PT = 9.0
# Average glyph advance as a share of the font size, and line height as a share of it.
_GLYPH_EM = 0.55
_LINE_EM = 1.2


# --- colours ----------------------------------------------------------------


def _luminance(hex_color: str) -> float:
    def channel(v: int) -> float:
        c = v / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    r, g, b = (int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * channel(r) + 0.7152 * channel(g) + 0.0722 * channel(b)


def _contrast(a: str, b: str) -> float:
    hi, lo = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


class Palette:
    """The colours a visual on this slide may use, taken from the deck's theme."""

    def __init__(self, deck: Deck, slide: Slide):
        theme = deck.theme_colors
        self.background = self._background(slide, theme)
        accents = [theme[k] for k in (f"accent{i}" for i in range(1, 7)) if theme.get(k)]
        readable = [c for c in accents if _contrast(c, self.background) >= 1.6]
        self.accents = readable or accents or ["0077FF"]
        self.text = "000000" if _luminance(self.background) > 0.4 else "FFFFFF"
        self.muted = self._blend(self.text, self.background, 0.35)

    @staticmethod
    def _background(slide: Slide, theme: dict[str, str]) -> str:
        color = slide.background
        if color is None:
            return "FFFFFF"
        if color.kind == "rgb" and color.rgb:
            return color.rgb
        return theme.get(color.theme_color or "", "FFFFFF")

    @staticmethod
    def _blend(a: str, b: str, share_b: float) -> str:
        parts = (
            round(int(a[i : i + 2], 16) * (1 - share_b) + int(b[i : i + 2], 16) * share_b)
            for i in (0, 2, 4)
        )
        return "".join(f"{p:02X}" for p in parts)

    def accent(self, i: int) -> str:
        return self.accents[i % len(self.accents)]

    def on(self, fill: str) -> str:
        """White or black, whichever reads better on `fill` (white unless it is clearly poor)."""
        white, black = _contrast("FFFFFF", fill), _contrast("000000", fill)
        return "FFFFFF" if white >= 3.0 or white >= black else "000000"


# --- text -------------------------------------------------------------------


def _fit_pt(texts: list[tuple[str, bool]], width: int, height: int, top_pt: float) -> float:
    """Largest size (<= `top_pt`) at which all `texts` (text, bold) fit one box."""
    pt = top_pt
    while pt > MIN_TEXT_PT:
        glyph = pt * _GLYPH_EM * _EMU_PT
        per_line = max(1.0, width / glyph)
        lines = sum(max(1, math.ceil(len(t) / per_line)) for t, _ in texts if t)
        if lines * pt * _LINE_EM * _EMU_PT <= height:
            return pt
        pt -= 0.5
    return MIN_TEXT_PT


def _text(
    shapes,
    x,
    y,
    w,
    h,
    text,
    pt,
    color,
    font,
    *,
    bold=False,
    align=PP_ALIGN.CENTER,
    anchor=MSO_ANCHOR.TOP,
):
    box = shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(max(1, int(w))), Emu(max(1, int(h))))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Emu(36_000)
    frame.margin_top = frame.margin_bottom = Emu(18_000)
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    if font:
        run.font.name = font
    return box


def _shape(shapes, kind, x, y, w, h, fill, *, rotation=0.0):
    shape = shapes.add_shape(
        getattr(MSO_SHAPE, kind), Emu(int(x)), Emu(int(y)), Emu(max(1, int(w))), Emu(max(1, int(h)))
    )
    shape.rotation = rotation
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _arrow_line(shapes, x1, y1, x2, y2, color, width_pt=1.75):
    connector = shapes.add_connector(1, Emu(int(x1)), Emu(int(y1)), Emu(int(x2)), Emu(int(y2)))
    connector.line.width = Pt(width_pt)
    connector.line.color.rgb = RGBColor.from_string(color)
    line = connector.line._get_or_add_ln()
    tail = line.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    line.append(tail)
    return connector


def _badge(group, item: DiagramItem, n: int, x, y, d, fill, palette: Palette, font, pt):
    """A filled circle holding the item's pictogram, or its number when it has none."""
    _shape(group.shapes, "OVAL", x, y, d, d, fill)
    on = palette.on(fill)
    if item.icon:
        pad = d * 0.2
        draw_icon(group.shapes, item.icon, int(x + pad), int(y + pad), int(d - 2 * pad), on, fill)
    else:
        _text(
            group.shapes,
            x,
            y,
            d,
            d,
            str(n),
            pt * 1.3,
            on,
            font,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )


# --- diagrams ---------------------------------------------------------------


def _draw_process(group, items, L, T, W, H, palette, font, top_pt):
    n = len(items)
    gap = W * 0.04
    cell = (W - gap * (n - 1)) / n
    d = min(cell * 0.7, H * 0.4)
    gap_y, label_h, detail_h = H * 0.04, H * 0.14, H * 0.26
    top = T + (H - (d + gap_y + label_h + detail_h)) / 2  # centred in its box
    pt = _fit_pt(
        [(i.label, True) for i in items] + [(i.detail, False) for i in items],
        int(cell),
        int(label_h + detail_h),
        top_pt,
    )
    for k, item in enumerate(items):
        x = L + k * (cell + gap)
        _badge(group, item, k + 1, x + (cell - d) / 2, top, d, palette.accent(k), palette, font, pt)
        if k < n - 1:
            _shape(
                group.shapes,
                "RIGHT_ARROW",
                x + cell + gap * 0.15,
                top + d * 0.35,
                gap * 0.7,
                d * 0.3,
                palette.muted,
            )
        _text(
            group.shapes,
            x,
            top + d + gap_y,
            cell,
            label_h,
            item.label,
            pt,
            palette.text,
            font,
            bold=True,
        )
        if item.detail:
            _text(
                group.shapes,
                x,
                top + d + gap_y + label_h,
                cell,
                detail_h,
                item.detail,
                pt * 0.9,
                palette.muted,
                font,
            )


def _draw_timeline(group, items, L, T, W, H, palette, font, top_pt):
    n = len(items)
    step = W / n
    mid = T + H / 2
    _shape(group.shapes, "RECTANGLE", L, mid - H * 0.008, W, H * 0.016, palette.muted)
    text_h = H * 0.42
    pt = _fit_pt(
        [(i.label, True) for i in items] + [(i.detail, False) for i in items],
        int(step * 1.6),
        int(text_h),
        top_pt,
    )
    d = H * 0.09
    for k, item in enumerate(items):
        cx = L + step * (k + 0.5)
        _shape(group.shapes, "OVAL", cx - d / 2, mid - d / 2, d, d, palette.accent(k))
        y = mid - d - text_h if k % 2 == 0 else mid + d
        w = step * 1.6
        x = min(max(cx - w / 2, L), L + W - w)
        _text(group.shapes, x, y, w, text_h * 0.3, item.label, pt, palette.text, font, bold=True)
        if item.detail:
            _text(
                group.shapes,
                x,
                y + text_h * 0.3,
                w,
                text_h * 0.7,
                item.detail,
                pt * 0.9,
                palette.muted,
                font,
            )


def _exit_distance(w: float, h: float, dx: float, dy: float) -> float:
    """How far along direction (dx, dy) a ray leaves a w x h box centred on its start."""
    tx = (w / 2) / abs(dx) if dx else math.inf
    ty = (h / 2) / abs(dy) if dy else math.inf
    return min(tx, ty)


def _draw_cycle(group, items, L, T, W, H, palette, font, top_pt):
    n = len(items)
    cx, cy = L + W / 2, T + H / 2
    node_w, node_h = W * 0.24, H * 0.2
    rx, ry = W / 2 - node_w / 2, H / 2 - node_h / 2
    centres = []
    for k in range(n):
        angle = -math.pi / 2 + 2 * math.pi * k / n
        centres.append((cx + rx * math.cos(angle), cy + ry * math.sin(angle)))
    for k in range(n):  # arrows first, so nodes sit on top of them
        (x1, y1), (x2, y2) = centres[k], centres[(k + 1) % n]
        dx, dy = x2 - x1, y2 - y1
        length = math.hypot(dx, dy) or 1.0
        out = _exit_distance(node_w, node_h, dx, dy) + 60_000 / length
        back = _exit_distance(node_w, node_h, -dx, -dy) + 60_000 / length
        if out + back < 1:
            _arrow_line(
                group.shapes,
                x1 + dx * out,
                y1 + dy * out,
                x2 - dx * back,
                y2 - dy * back,
                palette.muted,
            )
    pt = _fit_pt([(i.label, True) for i in items], int(node_w), int(node_h), top_pt)
    for k, ((x, y), item) in enumerate(zip(centres, items, strict=True)):
        fill = palette.accent(k)
        _shape(
            group.shapes, "ROUNDED_RECTANGLE", x - node_w / 2, y - node_h / 2, node_w, node_h, fill
        )
        _text(
            group.shapes,
            x - node_w / 2,
            y - node_h / 2,
            node_w,
            node_h,
            item.label,
            pt,
            palette.on(fill),
            font,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _draw_hierarchy(group, items, L, T, W, H, palette, font, top_pt):
    root, kids = items[0], items[1:]
    n = len(kids)
    root_w, root_h = min(W * 0.5, W), H * 0.24
    pt = _fit_pt([(i.label, True) for i in items], int(W / max(n, 1) * 0.9), int(H * 0.2), top_pt)
    rx = L + (W - root_w) / 2
    _shape(group.shapes, "ROUNDED_RECTANGLE", rx, T, root_w, root_h, palette.accent(0))
    _text(
        group.shapes,
        rx,
        T,
        root_w,
        root_h,
        root.label,
        pt * 1.1,
        palette.on(palette.accent(0)),
        font,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    if not kids:
        return
    gap = W * 0.03
    cell = (W - gap * (n - 1)) / n
    kid_top = T + H * 0.52
    kid_h = H * 0.3
    bar_y = T + root_h + (kid_top - T - root_h) / 2
    line = max(9525, H * 0.008)
    _shape(group.shapes, "RECTANGLE", L + cell / 2, bar_y, W - cell, line, palette.muted)
    _shape(
        group.shapes,
        "RECTANGLE",
        L + W / 2 - line / 2,
        T + root_h,
        line,
        bar_y - T - root_h,
        palette.muted,
    )
    for k, item in enumerate(kids):
        x = L + k * (cell + gap)
        _shape(
            group.shapes,
            "RECTANGLE",
            x + cell / 2 - line / 2,
            bar_y,
            line,
            kid_top - bar_y,
            palette.muted,
        )
        fill = palette.accent(k + 1)
        _shape(group.shapes, "ROUNDED_RECTANGLE", x, kid_top, cell, kid_h, fill)
        _text(
            group.shapes,
            x,
            kid_top,
            cell,
            kid_h,
            item.label,
            pt,
            palette.on(fill),
            font,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _draw_icons(group, items, L, T, W, H, palette, font, top_pt):
    n = len(items)
    cols = min(n, max(2, round(math.sqrt(n * W / max(H, 1)))))
    rows = math.ceil(n / cols)
    cw, ch = W / cols, H / rows
    d = min(cw * 0.42, ch * 0.5)
    label_h = ch - d - ch * 0.08
    pt = _fit_pt(
        [(i.label, True) for i in items] + [(i.detail, False) for i in items],
        int(cw * 0.9),
        int(label_h),
        top_pt,
    )
    for k, item in enumerate(items):
        x, y = L + (k % cols) * cw, T + (k // cols) * ch
        _badge(group, item, k + 1, x + (cw - d) / 2, y, d, palette.accent(k), palette, font, pt)
        _text(
            group.shapes,
            x,
            y + d + ch * 0.04,
            cw,
            label_h * 0.4,
            item.label,
            pt,
            palette.text,
            font,
            bold=True,
        )
        if item.detail:
            _text(
                group.shapes,
                x,
                y + d + ch * 0.04 + label_h * 0.4,
                cw,
                label_h * 0.6,
                item.detail,
                pt * 0.9,
                palette.muted,
                font,
            )


_DIAGRAMS = {
    "process": _draw_process,
    "timeline": _draw_timeline,
    "cycle": _draw_cycle,
    "hierarchy": _draw_hierarchy,
    "icons": _draw_icons,
}
MAX_ITEMS = {"process": 6, "timeline": 6, "cycle": 6, "hierarchy": 7, "icons": 8}


def draw_diagram(slide, shape: DiagramShape, palette: Palette) -> None:
    items = shape.items[: MAX_ITEMS[shape.diagram_type]]
    if not items:
        return
    group = slide.shapes.add_group_shape()
    _DIAGRAMS[shape.diagram_type](
        group,
        items,
        shape.left,
        shape.top,
        shape.width,
        shape.height,
        palette,
        shape.font_name,
        shape.font_size_pt or DEFAULT_TEXT_PT,
    )
    group.name = f"Диаграмма: {shape.diagram_type}"


# --- charts -----------------------------------------------------------------


def draw_chart(slide, shape: ChartShape, palette: Palette) -> None:
    if not shape.categories or not shape.series:
        return
    data = CategoryChartData()
    data.categories = shape.categories
    for series in shape.series:
        data.add_series(series.name, series.values)
    frame = slide.shapes.add_chart(
        _CHART_TYPES[shape.chart_type],
        Emu(shape.left),
        Emu(shape.top),
        Emu(shape.width),
        Emu(shape.height),
        data,
    )
    frame.name = f"График: {shape.chart_type}"
    chart = frame.chart
    chart.has_title = bool(shape.title)
    if shape.title:
        title = chart.chart_title.text_frame
        title.text = shape.title
        run = title.paragraphs[0].runs[0]
        run.font.size = Pt(shape.font_size_pt or DEFAULT_TEXT_PT)
        run.font.bold = False
        run.font.color.rgb = RGBColor.from_string(palette.text)
        chart.chart_title.include_in_layout = False
    chart.font.size = Pt(shape.font_size_pt or DEFAULT_TEXT_PT)
    chart.font.color.rgb = RGBColor.from_string(palette.text)
    if shape.font_name:
        chart.font.name = shape.font_name

    round_chart = shape.chart_type in ("pie", "doughnut")
    chart.has_legend = round_chart or len(shape.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    plot = chart.plots[0]
    if round_chart:
        plot.has_data_labels = True
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
        plot.data_labels.number_format = "0%"
        plot.data_labels.number_format_is_linked = False
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor.from_string(palette.accent(i))
        return

    for i, series in enumerate(plot.series):
        color = RGBColor.from_string(palette.accent(i))
        if shape.chart_type == "line":
            series.format.line.color.rgb = color
            series.format.line.width = Pt(2.5)
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.smooth = False
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
    if shape.chart_type != "line" and len(shape.series) == 1:
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt((shape.font_size_pt or DEFAULT_TEXT_PT) * 0.9)
    grid = RGBColor.from_string(Palette._blend(palette.text, palette.background, 0.85))
    chart.value_axis.major_gridlines.format.line.color.rgb = grid
    chart.value_axis.format.line.fill.background()
    for axis, title in (
        (chart.category_axis, shape.category_label),
        (chart.value_axis, shape.unit),
    ):
        if title:
            axis.has_title = True
            axis.axis_title.text_frame.text = title
            run = axis.axis_title.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt((shape.font_size_pt or DEFAULT_TEXT_PT) * 0.9)
            run.font.bold = False
            run.font.color.rgb = RGBColor.from_string(palette.muted)


def add_visuals(slide, slide_ir: Slide, deck: Deck) -> None:
    """Draw every chart and diagram of `slide_ir` onto the exported `slide`."""
    palette = Palette(deck, slide_ir)
    for shape in slide_ir.shapes:
        if isinstance(shape, ChartShape):
            draw_chart(slide, shape, palette)
        elif isinstance(shape, DiagramShape):
            draw_diagram(slide, shape, palette)
