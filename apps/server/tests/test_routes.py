"""The API surface stays intact across module splits.

uv run pytest apps/server
"""

from __future__ import annotations

from fastapi.testclient import TestClient
from server.main import app

client = TestClient(app)

EXPECTED = {
    ("get", "/health"),
    ("get", "/api/templates"),
    ("post", "/api/templates"),
    ("get", "/api/templates/{template_id}/inspect"),
    ("post", "/api/outline"),
    ("post", "/api/content"),
    ("post", "/api/audit"),
    ("post", "/api/audit/stream"),
    ("post", "/api/layout"),
    ("post", "/api/export"),
    ("post", "/api/export/pdf"),
    ("post", "/api/preview"),
    ("post", "/api/audit/deep"),
}


def test_every_public_route_is_registered():
    schema = client.get("/openapi.json").json()["paths"]
    have = {(method, path) for path, ops in schema.items() for method in ops}
    assert EXPECTED <= have


def test_health():
    assert client.get("/health").json() == {"status": "ok"}


def test_templates_are_listed_and_inspectable():
    templates = client.get("/api/templates").json()["templates"]
    assert templates
    first = templates[0]["id"]
    body = client.get(f"/api/templates/{first}/inspect").json()
    assert body["slides"] and body["deck"]["slides"]


def test_unknown_template_is_404():
    assert client.get("/api/templates/no-such-template/inspect").status_code == 404


def _minimal_deck() -> dict:
    return {
        "slide_width": 9_144_000,
        "slide_height": 6_858_000,
        "slides": [{"index": 0, "layout_name": "L", "shapes": []}],
    }


def test_deep_audit_is_501_without_a_renderer(monkeypatch):
    from export.render import RenderUnavailable

    def unavailable(_deck):
        raise RenderUnavailable("LibreOffice is not installed")

    monkeypatch.setattr("server.main._rendered_pages", unavailable)
    res = client.post("/api/audit/deep", json={"deck": _minimal_deck(), "brief": "тест"})
    assert res.status_code == 501


def test_deep_audit_returns_model_findings_when_rendered(monkeypatch):
    from audit import Finding

    monkeypatch.setattr("server.main._rendered_pages", lambda _deck: [b"fake-png-bytes"])
    monkeypatch.setattr(
        "server.main.run_model_checks",
        lambda deck, brief, images: [
            Finding(check="typo", kind="model", slide_index=0, message="found one")
        ],
    )
    res = client.post("/api/audit/deep", json={"deck": _minimal_deck(), "brief": "тест"})
    assert res.status_code == 200
    findings = res.json()["findings"]
    assert findings == [
        {
            "check": "typo",
            "kind": "model",
            "slide_index": 0,
            "shape_id": None,
            "message": "found one",
        }
    ]


def test_uploaded_template_is_private_not_in_the_shared_library(tmp_path):
    from pptx import Presentation
    from server import storage

    pptx_path = tmp_path / "my deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(pptx_path))

    with pptx_path.open("rb") as f:
        res = client.post(
            "/api/templates",
            files={
                "file": (
                    "my deck.pptx",
                    f,
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
    assert res.status_code == 200
    body = res.json()
    template_id = body["id"]
    assert template_id.startswith("upload:")

    # Never shows up in the shared library other sessions see.
    listed_ids = {t["id"] for t in client.get("/api/templates").json()["templates"]}
    assert template_id not in listed_ids

    # But the id the upload handed back still resolves, for its own uploader.
    inspect = client.get(f"/api/templates/{template_id}/inspect")
    assert inspect.status_code == 200
    assert inspect.json()["slides"]

    # Lives in the private upload store, not the shared evals/templates/ dir.
    stem = template_id.removeprefix("upload:")
    assert (storage.UPLOADED_TEMPLATES_DIR / f"{stem}.pptx").is_file()


def test_an_empty_template_library_gets_the_bundled_samples(tmp_path):
    from server.main import _ensure_sample_templates

    assert _ensure_sample_templates(tmp_path) is True
    assert len(list(tmp_path.glob("*.pptx"))) == 3
    # A library that already has templates is left alone.
    before = sorted(p.name for p in tmp_path.glob("*.pptx"))
    assert _ensure_sample_templates(tmp_path) is False
    assert sorted(p.name for p in tmp_path.glob("*.pptx")) == before


def test_deep_audit_uses_images_from_the_browser_without_any_renderer(monkeypatch):
    from audit import Finding

    def boom(_deck):
        raise AssertionError("must not render when images are supplied")

    seen = {}

    def fake_checks(deck, brief, images):
        seen["images"] = images
        return [Finding(check="typo", kind="model", slide_index=0, message="x")]

    monkeypatch.setattr("server.main._rendered_pages", boom)
    monkeypatch.setattr("server.main.run_model_checks", fake_checks)
    res = client.post(
        "/api/audit/deep",
        json={"deck": _minimal_deck(), "brief": "т", "images": ["data:image/png;base64,AAAA"]},
    )
    assert res.status_code == 200
    assert seen["images"] == {0: "data:image/png;base64,AAAA"}


def test_deep_audit_rejects_a_wrong_number_or_kind_of_images():
    deck = _minimal_deck()
    assert client.post("/api/audit/deep", json={"deck": deck, "images": []}).status_code == 400
    bad = client.post("/api/audit/deep", json={"deck": deck, "images": ["http://evil/x.png"]})
    assert bad.status_code == 400


def test_export_returns_a_pptx():
    res = client.post("/api/export", json=_minimal_deck())
    assert res.status_code == 200
    assert res.content[:2] == b"PK"  # a .pptx is a zip


def test_pdf_export_is_501_without_libreoffice(monkeypatch):
    from export.render import RenderUnavailable

    def unavailable(_pptx, _out_dir):
        raise RenderUnavailable("LibreOffice is not installed")

    monkeypatch.setattr("server.main.pptx_to_pdf", unavailable)
    assert client.post("/api/export/pdf", json=_minimal_deck()).status_code == 501


def test_pdf_export_returns_the_converted_file(monkeypatch):
    def convert(_pptx, out_dir):
        path = out_dir / "deck.pdf"
        path.write_bytes(b"%PDF-1.7 fake")
        return path

    monkeypatch.setattr("server.main.pptx_to_pdf", convert)
    res = client.post("/api/export/pdf", json=_minimal_deck())
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
